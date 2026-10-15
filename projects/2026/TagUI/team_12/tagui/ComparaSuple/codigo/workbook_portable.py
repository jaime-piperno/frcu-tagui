"""Genera el Excel con precios, estados y enlaces a la evidencia de cada consulta."""
from datetime import datetime, timezone
from math import ceil
from pathlib import Path, PurePosixPath
import re
from tempfile import NamedTemporaryFile
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties


STATUS = {
    'ok': 'Con stock + precio', 'revisar': 'Revisar variante',
    'sin_stock': 'Sin stock', 'no_encontrado': 'No encontrado en buscador',
    'marca_no_encontrada': 'Marca no encontrada en buscador',
    'sin_precio': 'Sin precio', 'error_consulta': 'Error de consulta',
    'error_extraccion': 'Cambio de página', 'bloqueado': 'Requiere intervención',
    'error_auth': 'Acceso no confirmado',
    'busqueda_incompleta': 'Búsqueda incompleta', 'error_url': 'URL inesperada',
    'pendiente': 'Pendiente', 'no_consultado': 'No consultado',
    'omitido_sin_sesion': 'Omitido sin sesión',
}
RED = 'A7242B'
DARK = '191919'
GREEN = 'DCF0DF'
AMBER = 'FFF2CC'
GRAY = 'F4F5F7'
MONEY = '"$" #,##0.00'


def _status(offer):
    if offer['status'] == 'bloqueado':
        reason = {'captcha': 'CAPTCHA', 'login_requerido': 'iniciar sesión',
                  'acceso_denegado': 'acceso denegado'}.get(offer.get('block_type'))
        return 'Requiere intervención' + (': ' + reason if reason else '')
    return STATUS.get(offer['status'], offer['status'])


def _safe(value):
    # El contenido de las tiendas es texto; sólo nuestras fórmulas pueden ejecutarse.
    return "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@', '\t', '\r', '\n')) else value


def _relative_path(value):
    """Los enlaces internos usan barras compatibles con Windows y Linux."""
    value = str(value or '').replace('\\', '/')
    path = PurePosixPath(value)
    if not value or path.is_absolute() or '..' in path.parts or ':' in value:
        return ''
    return path.as_posix()


def _link(cell, target):
    if target:
        cell.hyperlink = target
        cell.font = Font(name='Arial', size=10, color='0563C1', underline='single',
                         bold=isinstance(cell.value, (int, float)))


def _web_url(value):
    try:
        url = urlsplit(str(value or ''))
        return str(value) if url.scheme in ('https', 'http') and url.netloc else ''
    except ValueError:
        return ''


def _band(sheet, row, last_column, text, size=10):
    """Las categorías ocupan una fila propia, sin repetirlas en cada producto."""
    for cells in sheet.iter_rows(min_row=row, max_row=row, max_col=last_column):
        for cell in cells:
            cell.fill = PatternFill('solid', fgColor=DARK)
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
    cell = sheet.cell(row, 1, _safe(text))
    cell.font = Font(name='Arial', size=size, bold=True, color='FFFFFF')
    cell.alignment = Alignment(vertical='center', indent=1)


def _updated(data):
    """Muestra la fecha de las consultas, no la fecha de regeneración del Excel."""
    value = max((r['checked_at'] for r in data['results'] if r.get('checked_at')),
                default=data['created_at'])
    try:
        date = datetime.fromisoformat(value.replace('Z', '+00:00'))
        if date.tzinfo is None:
            date = date.replace(tzinfo=timezone.utc)
        return date.astimezone().strftime('%d/%m/%Y %H:%M')
    except (ValueError, TypeError):
        return str(value)


def _tabular(wb, name, headers, rows, widths):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([_safe(value) for value in row])
    ws.freeze_panes = 'C2'
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{max(1, ws.max_row)}'
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor=RED)
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 36
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 31
        for cell in row:
            cell.font = Font(name='Arial', size=10, color=DARK)
            cell.alignment = Alignment(vertical='center', wrap_text=True, indent=1)
            cell.fill = PatternFill('solid', fgColor='FFFFFF' if cell.row % 2 == 0 else GRAY)
    return ws


def export_workbook(data, target):
    """Recibe los resultados ya procesados y reemplaza el Excel sólo al terminar."""
    target = Path(target)
    all_providers = data['providers']
    selected = data.get('selected_providers') or [p['id'] for p in all_providers]
    # La portada muestra las tiendas de esta corrida. El detalle conserva todas.
    providers = [p for p in all_providers if p['id'] in selected]
    if not 1 <= len(providers) <= 16377:
        raise ValueError('El comparador requiere entre 1 y 16377 proveedores.')
    last_column = 5 + len(providers)
    last_letter = get_column_letter(last_column)
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Comparación'
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.showRowColHeaders = False
    sheet.sheet_view.zoomScale = 90
    sheet.sheet_properties.tabColor = RED
    sheet.freeze_panes = 'B7'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if len(providers) > 3 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:6'
    sheet.page_margins.left = sheet.page_margins.right = .25
    sheet.oddFooter.center.text = 'Página &P de &N'
    for column, width in enumerate([56, 15, 18, 19, 23] + [20] * len(providers), 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row_no in (3, 4, 5):
        sheet.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=last_column)
    sheet.row_dimensions[1].height = 12
    _band(sheet, 2, last_column, 'COMPARADOR DE PRECIOS', size=16)
    sheet['A2'].alignment = Alignment(horizontal='center', vertical='center')
    sheet.row_dimensions[2].height = 38
    confirmed = {r['product_id'] for r in data['results']
                 if r['provider'] in selected and r['status'] == 'ok' and r.get('price') is not None}
    sheet['A3'] = (f"Consultado: {_updated(data)} (hora local) · Productos: {len(data['products'])} · "
                   f"Con precio confirmado: {len(confirmed)} · Tiendas: {len(providers)}")
    sheet['A4'] = 'Importes en ARS, sin envío. Precio de venta: valor del catálogo. Menor precio: ofertas confirmadas de esta corrida.'
    sheet['A5'] = 'Clic en un precio de tienda para abrir la publicación. Consultas limitadas, variantes y evidencias: hoja Detalle.'
    for row_no in (3, 4, 5):
        sheet.cell(row_no, 1).font = Font(name='Arial', size=10, color='59616B')
        sheet.cell(row_no, 1).alignment = Alignment(vertical='center', wrap_text=True)
        sheet.row_dimensions[row_no].height = 28 if row_no == 4 else 23
    headers = ['PRODUCTO', 'COSTO', 'PRECIO DE VENTA', 'MENOR PRECIO', 'TIENDA DE REFERENCIA'] + [p['name'].upper() for p in providers]
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(6, col, _safe(header))
        cell.fill = PatternFill('solid', fgColor=RED)
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(right=Side(style='thin', color='E8BEC1'))
    sheet.row_dimensions[6].height = 36
    rows = {(r['product_id'], r['provider']): r for r in data['results']}
    details = []
    groups = {}
    for product in data['products']:
        category = str(product.get('category') or 'Sin categoría').strip() or 'Sin categoría'
        groups.setdefault(category, []).append(product)
    row_no = 6
    for category, products in groups.items():
        row_no += 1
        _band(sheet, row_no, last_column, f'{category.upper()} ({len(products)})')
        sheet.row_dimensions[row_no].height = 23
        first_product_row = row_no + 1
        for index, product in enumerate(products):
            row_no += 1
            offers = {}
            for provider in all_providers:
                offer = rows.get((product['id'], provider['id'])) or {'status': 'pendiente' if provider['id'] in selected else 'no_consultado'}
                price = offer.get('price') if offer['status'] == 'ok' else None
                offers[provider['id']] = offer
                refs = offer.get('evidence') or []
                # La última ficha abierta es la que confirmó o descartó el precio.
                fichas = [ref for ref in refs if re.search(r'-ficha-\d+\.json$', ref)]
                evidence = _relative_path((fichas or refs or [''])[-1])
                details.append([product['id'], product['name'], provider['name'], _status(offer),
                                price, offer.get('method') if price is not None else '', offer.get('title'), offer.get('variant_id'),
                                offer.get('url'), offer.get('checked_at'), offer.get('reason'), evidence])
            values = [product['name'], product.get('cost'), product.get('sale_price'), None, None]
            values += [offers[p['id']].get('price') if offers[p['id']]['status'] == 'ok' else
                       _status(offers[p['id']]) for p in providers]
            for col, value in enumerate(values, 1):
                cell = sheet.cell(row_no, col, _safe(value))
                cell.font = Font(name='Arial', size=10, color=DARK, bold=col in (3, 4))
                cell.fill = PatternFill('solid', fgColor='FFFFFF' if index % 2 == 0 else GRAY)
                cell.alignment = Alignment(horizontal='left' if col in (1, 5) else 'right',
                                           vertical='center', wrap_text=True, indent=1)
                cell.border = Border(bottom=Side(style='hair', color='DDE0E4'))
                if col in (2, 3, 4) or col >= 6:
                    cell.number_format = MONEY
                if col >= 6:
                    offer = offers[providers[col - 6]['id']]
                    _link(cell, _web_url(offer.get('url')))
                    if offer['status'] != 'ok':
                        color = 'FBE5E5' if offer['status'] == 'sin_stock' else AMBER
                        if offer['status'] in ('no_encontrado', 'marca_no_encontrada', 'no_consultado', 'pendiente'):
                            color = GRAY
                        cell.fill = PatternFill('solid', fgColor=color)
            # Las fórmulas son sólo D/E. Los nombres de tiendas y productos siguen siendo texto seguro.
            prices = f'F{row_no}:{last_letter}{row_no}'
            sheet.cell(row_no, 4, f'=IF(COUNT({prices})=0,"Sin referencia",MIN({prices}))')
            reference = '$F$6' if len(providers) == 1 else f'INDEX($F$6:${last_letter}$6,1,MATCH(D{row_no},{prices},0))'
            sheet.cell(row_no, 5, f'=IF(COUNT({prices})=0,"",{reference})')
            lines = max(ceil(len(str(value or '')) / (48 if col == 1 else 18))
                        for col, value in enumerate(values, 1))
            sheet.row_dimensions[row_no].height = max(36, 14 * lines + 8)
        for area, condition, color in (
            (f'C{first_product_row}:C{row_no}', f'ISNUMBER(C{first_product_row})', GREEN),
            (f'D{first_product_row}:D{row_no}', f'ISNUMBER(D{first_product_row})', GREEN),
            (f'D{first_product_row}:D{row_no}', f'NOT(ISNUMBER(D{first_product_row}))', AMBER),
            (f'F{first_product_row}:{last_letter}{row_no}', f'AND(ISNUMBER(F{first_product_row}),F{first_product_row}=$D{first_product_row})', GREEN),
        ):
            sheet.conditional_formatting.add(area, FormulaRule(formula=[condition], fill=PatternFill('solid', fgColor=color)))
    # Ordenar una tabla con categorías fusionadas rompería sus grupos; el filtro está en Detalle.
    sheet.print_area = f'A1:{last_letter}{row_no}'
    detail = _tabular(wb, 'Detalle', ['ID', 'Producto solicitado', 'Proveedor', 'Estado', 'Precio ARS', 'Forma de pago',
        'Variante encontrada', 'ID variante', 'Publicación', 'Consulta UTC', 'Observación', 'Evidencia'], details,
        [12, 58, 20, 27, 17, 22, 60, 18, 22, 28, 55, 20])
    for row in detail.iter_rows(min_row=2):
        if isinstance(row[4].value, (int, float)):
            row[4].number_format = MONEY
        url = _web_url(row[8].value)
        evidence = _relative_path(row[11].value)
        if url:
            row[8].value = 'Ver publicación'
            _link(row[8], url)
        if evidence:
            row[11].value = 'Ver evidencia'
            _link(row[11], evidence)
        lines = max(ceil(len(str(cell.value or '')) / max(10, detail.column_dimensions[cell.column_letter].width - 6)) for cell in row)
        detail.row_dimensions[row[0].row].height = max(34, 14 * lines + 8)
    wb.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)
    target.parent.mkdir(parents=True, exist_ok=True)
    # El temporal comparte carpeta: replace es atómico y no daña un Excel abierto.
    with NamedTemporaryFile(dir=target.parent, suffix='.xlsx', delete=False) as file:
        temporary = Path(file.name)
    try:
        wb.save(temporary)
        temporary.replace(target)
    finally:
        wb.close()
        temporary.unlink(missing_ok=True)
    return target
